"""Read-only comparison of an Excel/LibreOffice-recalculated copy to Python."""
import argparse
import json
from pathlib import Path
from openpyxl import load_workbook
from labs.day4.office_lab.sheets.student_excel import load_sample, calculate_wbs, grade_exam


def verify(directory: Path) -> dict:
    wbs = load_workbook(directory / "WBS_Gantt.xlsx", data_only=True)["WBS"]
    expected_wbs = calculate_wbs(load_sample("wbs_tasks.json"), "2026-09-25")
    for row, task in enumerate(expected_wbs, 9):
        assert wbs.cell(row, 8).value == task["working_days"], (row, "working_days")
        assert wbs.cell(row, 9).value == task["status"], (row, "status")
        assert wbs.cell(row, 10).value == "입력 완료", (row, "input_status")
    exam = load_workbook(directory / "Exam_Grading.xlsx", data_only=True)
    sheet = exam["채점"]
    expected_exam = grade_exam(load_sample("exam_sample.json"))
    for row, student in enumerate(expected_exam["students"], 10):
        assert sheet.cell(row, 44).value == (student["score"] if student["score"] is not None else "채점 보류"), (row, "score", sheet.cell(row, 44).value)
        assert sheet.cell(row, 45).value == (student["rank"] if student["rank"] is not None else "순위 제외"), (row, "rank")
        assert sheet.cell(row, 46).value == student["status"], (row, "status")
    for column, question in enumerate(expected_exam["questions"], 4):
        assert sheet.cell(41, column).value == question["correct"], (column, "correct")
        assert sheet.cell(42, column).value == question["denominator"], (column, "denominator")
        assert abs(sheet.cell(43, column).value - question["rate"]) < 1e-10, (column, "rate")
    for book in [exam, load_workbook(directory / "WBS_Gantt.xlsx", data_only=True)]:
        for tab in book:
            for row in tab:
                for cell in row:
                    assert cell.data_type != "e", (tab.title, cell.coordinate, cell.value)
    features = {}
    for filename in ["WBS_Gantt.xlsx", "Exam_Grading.xlsx"]:
        book = load_workbook(directory / filename)
        features[filename] = {sheet.title: {"conditional_format_ranges": len(sheet.conditional_formatting),
                                           "freeze_panes": str(sheet.freeze_panes)} for sheet in book}
    mutations = {}
    if (directory / "WBS_my_project.xlsx").exists():
        changed_wbs = load_workbook(directory / "WBS_my_project.xlsx", data_only=True)["WBS"]
        assert changed_wbs["I12"].value == "완료"
        assert changed_wbs["H5"].value == 1
        mutations["W04_progress_100_percent"] = {"status": "완료", "late_tasks": 1}
    if (directory / "Exam_my_class.xlsx").exists():
        changed_exam = load_workbook(directory / "Exam_my_class.xlsx", data_only=True)["채점"]
        assert changed_exam["B5"].value == 27
        assert changed_exam["AT35"].value == "채점 완료"
        assert changed_exam["AR35"].value == 24
        mutations["S026_invalid_answer_repaired"] = {"eligible_count": 27, "score": 24}
    return {"engine": "LibreOffice native recalculation", "wbs_tasks_checked": 14,
            "exam_students_checked": 28, "question_rates_checked": 40,
            "formula_errors": 0, "features": features, "native_input_changes": mutations}


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("directory", type=Path)
    args = parser.parse_args()
    print(json.dumps(verify(args.directory), ensure_ascii=False, indent=2))
