"""Student exercises: edit template inputs, retain Excel formulas, verify independently.

The delivered templates are authored with artifact-tool. This public-library
exercise uses openpyxl only for the student's requested input edits. openpyxl
does not calculate Excel formulas: open the saved copy in Excel/LibreOffice.
"""
from __future__ import annotations

import argparse
import csv
import json
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

SAMPLES = Path(__file__).resolve().parent


def _safe_path(workspace: Path, relative: str | Path) -> Path:
    """Resolve symlinks before checking the configured boundary."""
    root = Path(workspace).resolve()
    target = (root / relative).resolve()
    if not target.is_relative_to(root):
        raise ValueError("PATH_OUTSIDE_WORKSPACE")
    return target


def load_sample(name: str) -> Any:
    if name not in {"wbs_tasks.json", "exam_sample.json"}:
        raise ValueError("UNKNOWN_SAMPLE")
    return json.loads((SAMPLES / name).read_text(encoding="utf-8"))


def _date(value: str) -> date:
    return date.fromisoformat(value)


def validate_wbs(tasks: list[dict]) -> list[dict[str, str]]:
    """Return actionable input errors; zero progress is valid, blanks are not."""
    errors = []
    ids = [task.get("id") for task in tasks]
    by_id = {task.get("id"): task for task in tasks}
    for task in tasks:
        task_id = task.get("id", "")
        def add(code: str) -> None:
            errors.append({"id": task_id, "code": code})
        if any(not task.get(field) for field in ("id", "task", "role", "start", "end")):
            add("REQUIRED_INPUT_MISSING")
            continue
        if ids.count(task_id) != 1:
            add("DUPLICATE_TASK_ID")
        try:
            start, end = _date(task["start"]), _date(task["end"])
        except (ValueError, TypeError):
            add("INVALID_DATE")
            continue
        if end < start:
            add("END_BEFORE_START")
        progress = task.get("progress")
        if isinstance(progress, bool) or not isinstance(progress, (int, float)) or not 0 <= progress <= 1:
            add("INVALID_PROGRESS")
        predecessor = task.get("predecessor", "")
        if predecessor:
            if predecessor not in by_id or predecessor == task_id:
                add("INVALID_PREDECESSOR")
            else:
                try:
                    if start <= _date(by_id[predecessor]["end"]):
                        add("PREDECESSOR_DATE_CONFLICT")
                except (ValueError, TypeError, KeyError):
                    add("INVALID_PREDECESSOR_DATE")
            seen = {task_id}
            cursor = predecessor
            while cursor in by_id:
                if cursor in seen:
                    add("DEPENDENCY_CYCLE")
                    break
                seen.add(cursor)
                cursor = by_id[cursor].get("predecessor", "")
    return errors


def calculate_wbs(tasks: list[dict], as_of: str) -> list[dict]:
    """Independent NETWORKDAYS and status oracle; weekend-only calendar."""
    errors = validate_wbs(tasks)
    if errors:
        raise ValueError("WBS_INPUT_INVALID", errors)
    reference = _date(as_of)
    result = []
    for task in tasks:
        start, end = _date(task["start"]), _date(task["end"])
        working_days = sum((start + timedelta(days=i)).weekday() < 5 for i in range((end-start).days+1))
        status = ("완료" if task["progress"] == 1 else "지연" if end < reference
                  else "예정" if start > reference else "진행 중")
        result.append({**task, "working_days": working_days, "status": status})
    return result


def _valid_answer(value: Any) -> bool:
    return not isinstance(value, bool) and isinstance(value, (int, float)) and value in (1, 2, 3, 4)


def grade_exam(data: dict) -> dict:
    """Grade 40 questions. Invalid rows are excluded, blanks count as wrong."""
    key = data.get("key", [])
    students = data.get("students", [])
    if len(key) != 40 or not isinstance(students, list):
        raise ValueError("EXAM_SCHEMA_INVALID")
    key_valid = all(_valid_answer(x) for x in key)
    ids = [student.get("id") for student in students]
    results = []
    for student in students:
        answers = student.get("answers", [])
        if len(answers) != 40:
            raise ValueError("ANSWER_COUNT_INVALID", student.get("id"))
        blank = sum(value is None or value == "" for value in answers)
        invalid = sum(value not in (None, "") and not _valid_answer(value) for value in answers)
        status = "채점 완료"
        if not student.get("id") or ids.count(student["id"]) != 1:
            status = "ID 확인"
        elif student.get("attendance") == "결시":
            status = "결시"
        elif student.get("attendance") != "응시":
            status = "응시 상태 확인"
        elif not key_valid:
            status = "정답 확인"
        elif invalid:
            status = "입력 오류"
        elif blank:
            status = "미응답 포함"
        score = sum(_valid_answer(a) and a == k for a, k in zip(answers, key)) if status in {"채점 완료", "미응답 포함"} else None
        results.append({"id": student.get("id"), "status": status, "score": score,
                        "rank": None, "blank_count": blank, "invalid_count": invalid})
    scores = [item["score"] for item in results if item["score"] is not None]
    for item in results:
        if item["score"] is not None:
            item["rank"] = 1 + sum(score > item["score"] for score in scores)
    eligible = [student for student, result in zip(students, results) if result["score"] is not None]
    questions = [{"question": i+1, "correct": sum(student["answers"][i] == key[i] for student in eligible),
                  "denominator": len(eligible),
                  "rate": sum(student["answers"][i] == key[i] for student in eligible)/len(eligible) if eligible else None}
                 for i in range(40)]
    return {"students": results, "questions": questions, "eligible_count": len(eligible),
            "mean": sum(scores)/len(scores) if scores else None}


def _copy_template(workspace: Path, template: str, output: str):
    from openpyxl import load_workbook
    from openpyxl.workbook.properties import CalcProperties
    source = _safe_path(workspace, template)
    destination = _safe_path(workspace, output)
    if destination == source or destination.exists():
        raise FileExistsError("OUTPUT_ALREADY_EXISTS")
    book = load_workbook(source)
    book.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True, forceFullCalc=True, calcMode="auto")
    destination.parent.mkdir(parents=True, exist_ok=True)
    return book, destination


def _input_cell(sheet, row: int, column: int, value: Any) -> None:
    """Keep imported text literal; a task label must never become a formula."""
    cell = sheet.cell(row, column)
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"


def _csv_value(value: Any) -> Any:
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def write_wbs_copy(workspace: Path, tasks: list[dict], as_of: str,
                   output: str = "outputs/day4-student/WBS_my_project.xlsx") -> Path:
    """Save 14 tasks to a new template copy, never overwrite the reference."""
    if len(tasks) != 14:
        raise ValueError("TEMPLATE_REQUIRES_14_TASKS")
    calculate_wbs(tasks, as_of)
    first_date = min(_date(task["start"]) for task in tasks)
    if (max(_date(task["end"]) for task in tasks) - first_date).days >= 28:
        raise ValueError("GANTT_WINDOW_EXCEEDS_28_DAYS")
    book, destination = _copy_template(workspace, "outputs/day4-document-automation/WBS_Gantt.xlsx", output)
    sheet = book["WBS"]
    sheet["B4"] = datetime.combine(_date(as_of), datetime.min.time())
    sheet["B5"] = datetime.combine(first_date, datetime.min.time())
    for row, task in enumerate(tasks, 9):
        for column, field in enumerate(("id", "task", "role", "start", "end", "progress", "predecessor"), 1):
            value = task.get(field, "")
            _input_cell(sheet, row, column, datetime.combine(_date(value), datetime.min.time()) if field in {"start", "end"} else value)
    book.save(destination)
    return destination


def write_exam_copy(workspace: Path, data: dict,
                    output: str = "outputs/day4-student/Exam_my_class.xlsx") -> Path:
    """Save 28 synthetic rows. Invalid answers remain visible for repair practice."""
    grade_exam(data)
    if len(data["students"]) != 28:
        raise ValueError("TEMPLATE_REQUIRES_28_STUDENTS")
    book, destination = _copy_template(workspace, "outputs/day4-document-automation/Exam_Grading.xlsx", output)
    sheet = book["채점"]
    for column, answer in enumerate(data["key"], 4):
        _input_cell(sheet, 8, column, answer)
    for row, student in enumerate(data["students"], 10):
        _input_cell(sheet, row, 1, student["id"])
        _input_cell(sheet, row, 3, student["attendance"])
        for column, answer in enumerate(student["answers"], 4):
            _input_cell(sheet, row, column, answer)
    book.save(destination)
    return destination


def write_grade_csv(workspace: Path, data: dict, output: str = "outputs/day4-student/grading_report.csv") -> Path:
    """Independent calculated report, usable without Excel or API accounts."""
    destination = _safe_path(workspace, output)
    if destination.exists():
        raise FileExistsError("OUTPUT_ALREADY_EXISTS")
    report = grade_exam(data)
    destination.parent.mkdir(parents=True, exist_ok=True)
    with destination.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=["id", "status", "score", "rank", "blank_count", "invalid_count"])
        writer.writeheader()
        writer.writerows({key: _csv_value(value) for key, value in row.items()} for row in report["students"])
    return destination


def main() -> None:
    parser = argparse.ArgumentParser(description="4주차 Excel 입력 수정 실습")
    parser.add_argument("--workspace", type=Path, default=Path.cwd())
    parser.add_argument("--output-dir", default="outputs/day4-student")
    args = parser.parse_args()
    tasks, exam = load_sample("wbs_tasks.json"), load_sample("exam_sample.json")
    tasks[3]["progress"] = 1
    exam["students"][25]["answers"][8] = 1
    for result in [write_wbs_copy(args.workspace, tasks, "2026-09-25", f"{args.output_dir}/WBS_my_project.xlsx"),
                   write_exam_copy(args.workspace, exam, f"{args.output_dir}/Exam_my_class.xlsx"),
                   write_grade_csv(args.workspace, exam, f"{args.output_dir}/grading_report.csv")]:
        print(result)
    print("Excel 또는 LibreOffice에서 .xlsx 파일을 열면 수식이 재계산됩니다. CSV는 Python 독립 계산 결과입니다.")


if __name__ == "__main__":
    main()
