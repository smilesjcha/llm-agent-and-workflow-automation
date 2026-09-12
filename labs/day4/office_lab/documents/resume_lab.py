"""Create Word resume examples without inventing career facts.

Public dependencies: python-docx. No LLM, account, network, or secret needed.
All source and destination paths must stay under an explicit workspace.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path
import re
from typing import Any

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


class DocumentLabError(ValueError):
    """Expected, named validation failure; safe to show in a Notebook."""

    def __init__(self, code: str, detail: str = "") -> None:
        self.code = code
        self.detail = detail
        super().__init__(f"{code}: {detail}" if detail else code)


def workspace_path(workspace: Path | str, candidate: Path | str) -> Path:
    root = Path(workspace).resolve()
    path = Path(candidate)
    resolved = (path if path.is_absolute() else root / path).resolve()
    if not resolved.is_relative_to(root):
        raise DocumentLabError("PATH_OUTSIDE_WORKSPACE", str(candidate))
    return resolved


def read_fact_map(workspace: Path | str, source: Path | str) -> dict[str, Any]:
    path = workspace_path(workspace, source)
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise DocumentLabError("FACT_SOURCE_INVALID", path.name) from exc
    if not isinstance(data, dict) or data.get("schema_version") != 1 or not isinstance(data.get("roles"), list) or not data["roles"]:
        raise DocumentLabError("FACT_SCHEMA_INVALID")
    if not isinstance(data.get("title"), str) or not isinstance(data.get("summary"), str):
        raise DocumentLabError("FACT_SCHEMA_INVALID")
    ids: set[str] = set()
    for role in data["roles"]:
        if not isinstance(role, dict) or not isinstance(role.get("facts"), list):
            raise DocumentLabError("FACT_SCHEMA_INVALID")
        if not all(isinstance(role.get(key), str) for key in ("company", "role", "domain")):
            raise DocumentLabError("FACT_SCHEMA_INVALID")
        for fact in role.get("facts", []):
            if not isinstance(fact, dict):
                raise DocumentLabError("FACT_SCHEMA_INVALID")
            if not all(isinstance(fact.get(key), str) and fact[key] for key in ("id", "text", "rewrite")):
                raise DocumentLabError("FACT_SCHEMA_INVALID")
            if fact["id"] in ids:
                raise DocumentLabError("DUPLICATE_FACT_ID", fact["id"])
            ids.add(fact["id"])
    if not ids:
        raise DocumentLabError("FACT_SCHEMA_INVALID")
    return data


def numeric_tokens(text: str) -> set[str]:
    """Preserve units and comparison ranges, not merely isolated digits."""
    return set(re.findall(r"\d+(?:\.\d+)?(?:[-~]\d+(?:\.\d+)?)?\s*(?:만명|명|초|일|년|개월|%|원)?", text))


def audit_rewrite(facts: dict[str, Any], proposals: dict[str, str]) -> dict[str, Any]:
    """Reject unknown references and new numeric claims; semantics need a person."""
    known = {fact["id"]: fact for role in facts["roles"] for fact in role["facts"]}
    unknown = set(proposals) - set(known)
    if unknown:
        raise DocumentLabError("UNKNOWN_FACT_REFERENCE", ", ".join(sorted(unknown)))
    for fact_id, text in proposals.items():
        if not isinstance(text, str) or not text.strip():
            raise DocumentLabError("EMPTY_REWRITE", fact_id)
        unsupported = numeric_tokens(text) - numeric_tokens(known[fact_id]["text"])
        if unsupported:
            raise DocumentLabError("UNSUPPORTED_NUMERIC_CLAIM", f"{fact_id}: {', '.join(sorted(unsupported))}")
    return {
        "status": "REQUIRES_HUMAN_REVIEW",
        "checked_fact_ids": sorted(proposals),
        "numeric_guard": "PASS",
        "semantic_claims_verified": False,
        "source_verified": False,
        "pending_verification": [key for key in proposals if known[key].get("needs_verification")],
    }


def _style_document(document: Any) -> None:
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.7)
    section.left_margin = section.right_margin = Cm(1.9)
    sizes = {"Normal": 10.5, "Title": 23, "Subtitle": 11, "Heading 1": 14, "Heading 2": 11.5, "List Bullet": 10.5}
    for name, size in sizes.items():
        style = document.styles[name]
        style.font.name = "NanumGothic"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for key in ("asciiTheme", "eastAsiaTheme", "hAnsiTheme", "cstheme"):
            fonts.attrib.pop(qn("w:" + key), None)
        fonts.set(qn("w:eastAsia"), "NanumGothic")
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.line_spacing = 1.16
    document.styles["Title"].paragraph_format.space_after = Pt(8)
    document.styles["Heading 1"].paragraph_format.space_before = Pt(12)
    document.styles["Heading 2"].paragraph_format.space_before = Pt(7)
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    # The stock Word template can carry a blue title border through inheritance.
    for element in document.styles.element.xpath(".//w:pBdr"):
        element.getparent().remove(element)


def build_resume(
    workspace: Path | str,
    source: Path | str,
    output: Path | str,
    *,
    variant: str = "after",
    proposals: dict[str, str] | None = None,
) -> dict[str, Any]:
    """Create a new document. Existing files are never silently overwritten."""
    if variant not in {"before", "after"}:
        raise DocumentLabError("INVALID_VARIANT", variant)
    facts = read_fact_map(workspace, source)
    destination = workspace_path(workspace, output)
    if destination.suffix.lower() != ".docx":
        raise DocumentLabError("INVALID_OUTPUT_EXTENSION")
    if destination.exists():
        raise DocumentLabError("OUTPUT_EXISTS", destination.name)
    selected = {fact["id"]: fact["rewrite"] for role in facts["roles"] for fact in role["facts"]}
    selected.update(proposals or {})
    report = audit_rewrite(facts, selected)
    doc = Document()
    _style_document(doc)
    doc.add_paragraph(facts["title"], style="Title")
    if variant == "before":
        doc.add_paragraph("저는 여러 산업에서 AI 업무를 해왔고, 현재 이커머스에서 Agent 관련 서비스의 PM 업무를 담당하고 있습니다. 과거에는 교육, 의료, 금융 분야에서 AI Engineer로 서비스를 만들고 운영했습니다.")
        doc.add_paragraph("업무 경력", style="Heading 1")
        for role in facts["roles"]:
            doc.add_paragraph(role["company"] + " " + role["role"], style="Heading 2")
            doc.add_paragraph("제가 진행했던 업무는 " + ". 또한 ".join(f["text"] for f in role["facts"]) + "입니다.")
    else:
        doc.add_paragraph(facts["summary"])
        doc.add_paragraph("현재 Role", style="Heading 1")
        doc.add_paragraph(facts["roles"][0]["company"] + "  " + facts["roles"][0]["role"], style="Heading 2")
        for fact in facts["roles"][0]["facts"]:
            doc.add_paragraph(selected[fact["id"]], style="List Bullet")
        doc.add_paragraph("산업별 경력", style="Heading 1")
        for role in facts["roles"][1:]:
            doc.add_paragraph(role["company"] + "  " + role["domain"], style="Heading 2")
            doc.add_paragraph(role["role"])
            for fact in role["facts"]:
                suffix = " (수치 확인 필요)" if fact.get("needs_verification") else ""
                doc.add_paragraph(selected[fact["id"]] + suffix, style="List Bullet")
    doc.add_paragraph("강의 및 멘토링", style="Heading 1")
    for item in facts["teaching"]:
        doc.add_paragraph(item, style="List Bullet" if variant == "after" else "Normal")
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.font.name = "NanumGothic"
            run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "NanumGothic")
    destination.parent.mkdir(parents=True, exist_ok=True)
    # Exclusive creation preserves the original even if another process wins the race.
    with destination.open("xb") as file:
        doc.save(file)
    report.update({"path": str(destination), "variant": variant, "paragraph_count": len(doc.paragraphs)})
    return report


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workspace", type=Path, default=Path.cwd())
    parser.add_argument("--source", default="labs/day4/office_lab/documents/career_facts.json")
    parser.add_argument("--output", required=True)
    parser.add_argument("--variant", choices=("before", "after"), default="after")
    args = parser.parse_args()
    try:
        print(json.dumps(build_resume(args.workspace, args.source, args.output, variant=args.variant), ensure_ascii=False, indent=2))
    except DocumentLabError as exc:
        print(json.dumps({"status": "ERROR", "error": exc.code, "detail": exc.detail}, ensure_ascii=False))
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
