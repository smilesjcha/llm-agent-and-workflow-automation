from pathlib import Path

from docx import Document
import pytest

from labs.day4.office_lab.documents.resume_lab import (
    DocumentLabError,
    audit_rewrite,
    build_resume,
    read_fact_map,
    workspace_path,
)
from labs.day4.office_lab.presentations.brief_data import prepare_brief

ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path("labs/day4/office_lab/documents/career_facts.json")


def test_word_before_after_preserve_facts(tmp_path):
    (tmp_path / "facts.json").write_bytes((ROOT / SOURCE).read_bytes())
    for variant in ("before", "after"):
        result = build_resume(tmp_path, "facts.json", f"{variant}.docx", variant=variant)
        doc = Document(result["path"])
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "크레버스" in text and "ABACUS" in text
        assert "7만명" in text and "3-5일" in text and "10초" in text
        assert doc.paragraphs[0].style.name == "Title"
        assert result["status"] == "REQUIRES_HUMAN_REVIEW"
        assert not result["semantic_claims_verified"]
    assert any(p.style.name == "List Bullet" for p in Document(tmp_path / "after.docx").paragraphs)


def test_new_numeric_claim_rejected():
    facts = read_fact_map(ROOT, SOURCE)
    with pytest.raises(DocumentLabError, match="UNSUPPORTED_NUMERIC_CLAIM"):
        audit_rewrite(facts, {"education-latency": "평가 시간을 1초로 단축"})


def test_unknown_fact_reference_rejected():
    with pytest.raises(DocumentLabError, match="UNKNOWN_FACT_REFERENCE"):
        audit_rewrite(read_fact_map(ROOT, SOURCE), {"new-revenue": "매출 개선"})


def test_original_is_not_overwritten(tmp_path):
    (tmp_path / "facts.json").write_bytes((ROOT / SOURCE).read_bytes())
    output = tmp_path / "original.docx"
    output.write_bytes(b"original")
    with pytest.raises(DocumentLabError, match="OUTPUT_EXISTS"):
        build_resume(tmp_path, "facts.json", output)
    assert output.read_bytes() == b"original"


@pytest.mark.parametrize("path", ["../outside.docx", "/etc/passwd"])
def test_paths_outside_workspace_rejected(tmp_path, path):
    with pytest.raises(DocumentLabError, match="PATH_OUTSIDE_WORKSPACE"):
        workspace_path(tmp_path, path)


def test_symlink_escape_rejected(tmp_path):
    (tmp_path / "escape").symlink_to(tmp_path.parent, target_is_directory=True)
    with pytest.raises(DocumentLabError, match="PATH_OUTSIDE_WORKSPACE"):
        workspace_path(tmp_path, "escape/outside.docx")


def test_source_schema_error_is_named(tmp_path):
    (tmp_path / "bad.json").write_text('{"schema_version": 99}', encoding="utf-8")
    with pytest.raises(DocumentLabError, match="FACT_SCHEMA_INVALID"):
        read_fact_map(tmp_path, "bad.json")


def test_non_word_output_rejected(tmp_path):
    (tmp_path / "facts.json").write_bytes((ROOT / SOURCE).read_bytes())
    with pytest.raises(DocumentLabError, match="INVALID_OUTPUT_EXTENSION"):
        build_resume(tmp_path, "facts.json", "resume.txt")


def test_ppt_source_preserves_wbs_and_review_counts():
    brief = prepare_brief(ROOT)
    assert len(brief["slides"]) == 7
    assert brief["sources"][0].endswith("wbs_tasks.json")
    assert len(brief["wbs_snapshot"]) == 14
    assert sum(int(row[1]) for row in brief["slides"][1]["table"][1:]) == 14
    assert len(brief["slides"][4]["table"]) == 3


def test_ppt_source_uses_edited_notebook_input():
    brief = prepare_brief(ROOT)
    edited = brief["wbs_snapshot"]
    edited[3]["progress"] = 1
    updated = prepare_brief(ROOT, tasks=edited)
    assert updated["slides"][2]["table"][4][-2:] == ["완료", "100%"]
    assert updated["sources"][0] == "Notebook WBS input snapshot"


def test_ppt_never_silently_omits_extra_tasks():
    tasks = prepare_brief(ROOT)["wbs_snapshot"]
    with pytest.raises(DocumentLabError, match="BRIEF_TASK_COUNT"):
        prepare_brief(ROOT, tasks=tasks[:7])
