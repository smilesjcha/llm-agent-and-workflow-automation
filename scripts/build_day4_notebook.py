"""Build the student Day 4 notebook from executable, reviewable lesson cells."""

from pathlib import Path
import argparse
import json
import sys
import tempfile
import textwrap
import nbformat as nbf

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "materials/day4/day4_pr_document_automation.ipynb"


def build():
    cells = []
    def md(source):
        cells.append(nbf.v4.new_markdown_cell(textwrap.dedent(source).strip()))
    def code(source):
        cells.append(nbf.v4.new_code_cell(textwrap.dedent(source).strip()))

    md("""
    # 4주차 GitHub PR 리뷰와 문서 자동화

    변경 코드를 확인하고, 테스트로 수정 결과를 검증합니다. 같은 원칙을 Excel 일정·채점, Word 이력서, PPT 보고서에 적용합니다.

    **실제로 만드는 파일:** 수정한 Python 코드, 테스트, 리뷰 Markdown·HTML, WBS Excel, 채점 Excel·CSV, Word 이력서, PPT 보고서, 결과 목록 HTML.

    기본 실행은 공개 합성 데이터와 제공한 양식만 사용합니다. 실제 GitHub 게시, Google Sheets 수정, 이메일 발송, LLM 호출은 자동으로 하지 않습니다. Codex 대화 과제는 원하는 경우 본인의 로그인 환경에서 직접 수행합니다.
    """)
    md("""
    ## 실행 준비

    VS Code에서 이 Notebook을 열고 `.venv` Python 커널을 선택합니다. 첫 설치에는 인터넷이 필요할 수 있습니다. 아래 설치 셀은 필요한 Python 패키지가 없을 때만 실행합니다.

    원본 자료를 덮어쓰지 않도록 실행할 때마다 새 결과 폴더를 만듭니다. `APPLY_REFERENCE_FIX=False`가 기본입니다. 학생 코드를 자동으로 고치지 않습니다. 의도한 테스트 실패는 학습 결과로 표시하며 전체 실행을 중단하지 않습니다.
    """)
    code("""
    from pathlib import Path
    import sys, os, uuid, json, shutil, subprocess, importlib.util
    from datetime import datetime, date, timedelta
    from copy import deepcopy

    candidates = [Path.cwd(), *Path.cwd().parents]
    ROOT = next((p for p in candidates if (p / 'labs/day4/pr_review_lab/service.py').is_file()), None)
    if ROOT is None:
        raise RuntimeError('저장소 전체를 받았는지 확인하고 해당 폴더에서 Notebook을 열어 주세요.')
    if str(ROOT) not in sys.path:
        sys.path.insert(0, str(ROOT))
    run_base = (ROOT / 'output/day4-notebook-runs').resolve()
    if not run_base.is_relative_to(ROOT):
        raise ValueError('PATH_OUTSIDE_WORKSPACE')
    RUN = run_base / (datetime.now().strftime('%Y%m%d_%H%M%S') + '_' + uuid.uuid4().hex[:6])
    RUN.mkdir(parents=True, exist_ok=False)
    APPLY_REFERENCE_FIX = False
    ALLOW_PUBLIC_GITHUB_READ = False
    print('Python:', sys.version.split()[0])
    print('결과 폴더:', RUN.relative_to(ROOT))
    print('학생 코드 자동 수정:', APPLY_REFERENCE_FIX)
    """)
    code("""
    # Notebook에서 직접 설치할 때도 현재 커널의 Python을 사용합니다.
    required = {'pytest': 'pytest', 'openpyxl': 'openpyxl', 'docx': 'python-docx', 'nbformat': 'nbformat'}
    missing = [package for module, package in required.items() if importlib.util.find_spec(module) is None]
    if missing:
        requirements = ROOT / 'requirements-day4.txt'
        command = [sys.executable, '-m', 'pip', 'install', '-r', str(requirements)]
        result = subprocess.run(command, cwd=ROOT, capture_output=True, text=True)
        if result.returncode:
            raise RuntimeError('패키지 설치 실패. 현재 커널과 인터넷 연결을 확인한 뒤 터미널에서 requirements-day4.txt를 설치하세요.')
        print('설치 완료:', ', '.join(missing))
    else:
        print('필수 Python 패키지 준비 완료')
    # 직접 실행 대안: %pip install -r ../../requirements-day4.txt
    """)
    code("""
    import html
    from IPython.display import display, HTML, Markdown, FileLink

    def show_rows(rows, columns=None):
        if not rows:
            print('표시할 행 없음')
            return
        columns = columns or list(rows[0])
        header = ''.join('<th>' + html.escape(str(key)) + '</th>' for key in columns)
        body = ''.join('<tr>' + ''.join('<td>' + html.escape(str(row.get(key, ''))) + '</td>' for key in columns) + '</tr>' for row in rows)
        display(HTML('<style>.day4-table{border-collapse:collapse;font-size:15px}.day4-table td,.day4-table th{border:1px solid #ccc;padding:8px 12px;text-align:left}.day4-table th{background:#111;color:white}</style><table class="day4-table"><tr>' + header + '</tr>' + body + '</table>'))

    def file_link(path):
        path = Path(path).resolve()
        if not path.is_relative_to(ROOT):
            raise ValueError('PATH_OUTSIDE_WORKSPACE')
        display(FileLink(os.path.relpath(path, Path.cwd())))

    def expected_failure(label, action, expected_code):
        try:
            action()
        except (ValueError, FileExistsError) as error:
            actual = getattr(error, 'code', str(error.args[0]))
            if actual != expected_code:
                raise
            print(label + ': ' + actual + ' (예상한 차단)')
        else:
            raise AssertionError(label + ': 차단되어야 할 입력이 통과했습니다.')

    artifacts = {}
    """)
    md("""
    ## 오늘의 시간표

    | 차시 | 시간 | 파일 작업 |
    |---|---|---|
    | 1 | 09:00–09:50 | PR 입력·diff·대상 검사 |
    | 2 | 09:50–10:40 | 코드 수정·리뷰 문서·사람 확인 |
    | 3 | 10:40–11:30 | pytest·CI YAML·중복과 버전 검사 |
    | 쉬는 시간 | 11:30–12:00 | 30분 |
    | 점심시간 | 12:00–13:00 | 60분 |
    | 4 | 13:00–13:50 | WBS·Gantt Chart |
    | 5 | 13:50–14:40 | 채점·조건부서식·통계 |
    | 쉬는 시간 | 14:40–15:00 | 20분 |
    | 6 | 15:00–15:50 | Word 이력서 개선 |
    | 7 | 15:50–16:40 | 데이터 기반 PPT 생성 |
    | 8 | 16:40–17:30 | 결과 연결·다음 주 미니 프로젝트 |
    | 쉬는 시간·Q&A | 17:30–18:00 | 30분 |
    """)

    md("""
    ## 1차시 대상 PR과 변경 코드

    PR 번호는 변경 제안의 주소이고, commit SHA는 그때의 코드 버전입니다. 리뷰할 코드가 바뀌면 같은 PR 안에서도 새 검토가 필요합니다. 먼저 데이터와 파일을 읽고, 목적지가 일치하는지 확인합니다.
    """)
    code("""
    from labs.day4.pr_review_lab import (
        load_fixture, validate_snapshot, prepare_exercise, check_exercise,
        build_review, approve_preview, preview_publication, render_outputs,
        fetch_public_pr, LabError,
    )
    snapshot = load_fixture()
    show_rows([{key: snapshot[key] for key in ('source', 'repo', 'number', 'head_sha')}])
    """)
    md("""
    ### 변경 줄

    `-`는 이전 코드, `+`는 새 코드입니다. 쿠폰을 추가한 반환식에서 초과 쿠폰과 잘못된 입력을 처리하는지 살펴봅니다. 이 PR은 수업용 합성 데이터이며 실제 GitHub 링크가 아닙니다.
    """)
    code("""
    print(snapshot['files'][0]['path'])
    print(snapshot['files'][0]['patch'])
    """)
    code("""
    def my_target_matches(pr, repo, number):
        # 직접 수정 구간: PR 대상 확인 함수를 한 줄씩 읽고 조건을 추가해 보세요.
        return pr['repo'].lower() == repo.lower() and pr['number'] == number

    assert my_target_matches(snapshot, snapshot['repo'], 42)
    assert not my_target_matches(snapshot, 'another/repository', 42)
    validate_snapshot(snapshot, expected_repo=snapshot['repo'], expected_pr=42)
    bad_target = deepcopy(snapshot)
    bad_target['head_sha'] = '최신버전'
    expected_failure('SHA 형식 검사', lambda: validate_snapshot(bad_target, expected_repo=snapshot['repo'], expected_pr=42), 'INVALID_SHA')
    """)
    md("""
    ### 내 실습 폴더

    원본 fixture를 보존하고 새 폴더에서 수정합니다. `checkout.py`는 기능 코드, `test_checkout.py`는 기대 동작입니다. 테스트를 약하게 바꾸는 대신 기능 코드를 고칩니다.
    """)
    code("""
    exercise = prepare_exercise(ROOT, RUN / 'pr/my-exercise')
    EXERCISE = Path(exercise['exercise_dir'])
    print('수정할 파일:', EXERCISE / 'checkout.py')
    print((EXERCISE / 'checkout.py').read_text(encoding='utf-8'))
    """)
    code("""
    baseline_test = check_exercise(ROOT, EXERCISE)
    print(baseline_test['output'])
    assert baseline_test['status'] == 'FAILED'
    print('의도된 시작 상태: 정상 1개 통과, 초과 쿠폰·잘못된 입력 4개 실패')
    """)
    md("""
    ### Codex 작업

    > 위에 출력된 내 실습 폴더의 checkout.py와 test_checkout.py를 읽어 줘. 실패하는 입력과 사용자 영향을 먼저 설명해 줘. 아직 파일을 수정하거나 GitHub에 게시하지 마. repo·PR 번호·commit SHA를 각각 왜 확인하는지도 내 코드와 연결해 설명해 줘.
    """)

    md("""
    ## 2차시 리뷰 초안과 코드 수정

    좋은 리뷰에는 변경 줄, 재현 입력, 사용자 영향, 최소 수정, 확인할 테스트가 있습니다. AI가 작성한 의견과 사람이 채택한 의견을 구분하고, 승인한 대상과 내용이 바뀌면 다시 확인합니다.
    """)
    code("""
    review = build_review(snapshot)
    display(Markdown(review['markdown']))
    print('출처:', review['provider'], '— 사람이 작성한 수업용 리뷰 예시이며 실시간 모델 출력이 아닙니다.')
    """)
    md("""
    ### 제품 규칙과 Role

    여기서는 가격·쿠폰 0 이상, 수량 1 이상, 최종 결제액 0 이상을 제품 규칙으로 정했습니다. 실제 서비스라면 소수 수량·부가세·반품 정책 등이 달라질 수 있으므로 먼저 규칙을 확인해야 합니다.
    """)
    code("""
    policy_path = ROOT / 'labs/day4/pr_review_lab/review_policy.md'
    display(Markdown(policy_path.read_text(encoding='utf-8')))
    """)
    md("""
    ### 선택 실습 Codex CLI의 실제 리뷰

    아래 셀의 기본값은 False입니다. 본인의 CLI 로그인·모델 접근·사용량을 확인한 뒤 True로 선택할 때만 실제 모델을 호출합니다. 공개 예시 세 파일을 별도 작업 폴더에 복사하고 read-only 모드로 리뷰를 요청합니다. 실패하면 fixture 결과로 바꾸지 않습니다.

    모델 원문의 줄 번호도 직접 확인합니다. 강사 실행 사례에서도 결함 설명은 맞았지만 반환식의 줄 번호를 틀린 경우가 있었습니다. 이 셀은 GitHub에 게시하거나 학생 코드를 수정하지 않습니다. read-only는 쓰기 제한이며, OS 수준의 세 파일 전용 읽기 격리를 보장하는 기능은 아닙니다.
    """)
    code("""
    from labs.day4.pr_review_lab.codex_review import run_codex_review
    RUN_CODEX_LIVE = False
    if RUN_CODEX_LIVE:
        try:
            codex_live = run_codex_review(ROOT, RUN/'pr/codex-review.raw.md', live=True, model='gpt-5.6-sol')
            display(Markdown(codex_live['markdown']))
            artifacts['Codex 실제 리뷰 원문'] = Path(codex_live['path'])
            lines = (ROOT/'labs/day4/pr_review_lab/fixtures/checkout.py').read_text(encoding='utf-8').splitlines()
            print('실제 파일 줄 번호:')
            print('\\n'.join(f'{i+1}: {line}' for i,line in enumerate(lines)))
            print('원문을 자동 교정하지 않습니다. 위치·재현 입력·제품 규칙을 직접 검증하세요.')
        except LabError as error:
            print(error.code, str(error))
            print('실제 모델 실행 미완료. fixture로 대체하지 않았습니다.')
    else:
        print('Codex 실제 호출 미실행: RUN_CODEX_LIVE=False. 위 리뷰는 수업용 fixture입니다.')
    """)
    md("""
    ### 직접 수정 구간

    위에서 출력된 **내 실습 폴더의 checkout.py**를 VS Code에서 엽니다. 잘못된 입력이면 `ValueError`를 발생시키고, 정상 계산 결과는 `max(0, ...)`으로 제한합니다. 다음 셀을 다시 실행하면 현재 수정본을 테스트합니다.

    처음부터 Run All을 하면 이 단계는 여전히 실패할 수 있습니다. 자동 수정은 기본으로 꺼져 있습니다. 참고 구현은 별도 파일이며 다음 차시에서 따로 테스트합니다.
    """)
    code("""
    if APPLY_REFERENCE_FIX:
        # 사용자가 직접 True로 선택한 경우에만 참고 구현을 학생 파일에 적용합니다.
        shutil.copyfile(EXERCISE / 'reference_solution.py', EXERCISE / 'checkout.py')
        print('참고 구현 적용 모드: 학생의 직접 수정 결과가 아닙니다.')
    current_test = check_exercise(ROOT, EXERCISE)
    print(current_test['output'])
    """)
    code("""
    HUMAN_CONFIRMED = False  # review['markdown']을 읽은 사람이 직접 True로 선택
    if HUMAN_CONFIRMED:
        approval = approve_preview(snapshot, review, approved=True, reviewer='수강생')
        publication = preview_publication(snapshot, review, approval, current_head_sha=snapshot['head_sha'])
        print(publication['status'], '실제 원격 게시:', publication['remote_write_performed'])
    else:
        approval, publication = None, None
        print('사람 확인 대기: 리뷰 초안만 저장합니다.')
    """)
    code("""
    pr_files = render_outputs(ROOT, RUN / 'pr/report', snapshot, review,
                              preview=publication, test_result=current_test)
    artifacts['PR 리뷰 화면'] = Path(pr_files['review.html'])
    artifacts['리뷰 Markdown'] = Path(pr_files['review.md'])
    file_link(pr_files['review.html'])
    file_link(pr_files['review.md'])
    """)
    md("""
    ### Codex 작업

    > 내가 수정한 checkout.py와 테스트 결과를 리뷰해 줘. 정상 쿠폰 계산을 유지했는지, 잘못된 입력을 명확히 거부하는지, 테스트를 약하게 바꾸지 않았는지 확인해 줘. 중요한 의견만 파일·줄·입력 예시와 함께 작성해 줘. merge와 게시 여부는 내가 결정할게.
    """)

    md("""
    ## 3차시 테스트와 CI

    CI는 같은 검사를 새 코드마다 반복하는 실행 환경입니다. 테스트 통과는 정해 둔 예제를 만족했다는 뜻이며, 제품 정책이나 표현의 적절성까지 보증하지 않습니다. 테스트 권한과 GitHub 게시 권한도 분리합니다.
    """)
    code("""
    reference_info = prepare_exercise(ROOT, RUN / 'pr/reference-only')
    REFERENCE = Path(reference_info['exercise_dir'])
    shutil.copyfile(REFERENCE / 'reference_solution.py', REFERENCE / 'checkout.py')
    reference_test = check_exercise(ROOT, REFERENCE)
    print('별도 참고 구현 검증. 내 실습 폴더의 파일은 수정하지 않았습니다.')
    print(reference_test['output'])
    assert reference_test['status'] == 'PASSED'
    show_rows([{'검사 대상': '수정 전 코드', '결과': baseline_test['status']},
               {'검사 대상': '현재 내 코드', '결과': current_test['status']},
               {'검사 대상': '별도 참고 구현', '결과': reference_test['status']}])
    """)
    md("""
    ### 승인 계약의 실패 테스트

    아래의 승인 객체는 **테스트용 합성 입력**입니다. 사람의 실제 승인이라고 저장하거나 게시하지 않습니다. 새 commit과 같은 리뷰 재실행을 각각 차단하는지 확인합니다.
    """)
    code("""
    simulated_approval = approve_preview(snapshot, review, approved=True, reviewer='fixture-test-only')
    expected_failure('새 commit', lambda: preview_publication(snapshot, review, simulated_approval,
        current_head_sha='3' * 40), 'STALE_HEAD_SHA')
    marker = '<!-- day4-review:' + review['review_sha256'] + ' -->'
    expected_failure('같은 리뷰 재실행', lambda: preview_publication(snapshot, review, simulated_approval,
        current_head_sha=snapshot['head_sha'], existing_reviews=[{'body': marker}]), 'DUPLICATE_REVIEW')
    expected_failure('승인 누락', lambda: preview_publication(snapshot, review, None,
        current_head_sha=snapshot['head_sha']), 'HUMAN_APPROVAL_REQUIRED')
    """)
    md("""
    ### GitHub Actions YAML

    제공 파일은 자동 실행되지 않는 템플릿입니다. 본인 저장소에 복사하기 전에 trigger, 작업 폴더, Python 버전, 실행 명령을 확인합니다. `pull_request`와 `contents: read`를 사용하고 secrets와 원격 게시 단계는 넣지 않았습니다.
    """)
    code("""
    workflow_source = ROOT / 'labs/day4/pr_review_lab/templates/workflow.yml'
    workflow_draft = RUN / 'pr/day4-checkout-tests.yml'
    workflow_draft.write_text(workflow_source.read_text(encoding='utf-8'), encoding='utf-8')
    print(workflow_draft.read_text(encoding='utf-8'))
    artifacts['CI 설정 초안'] = workflow_draft
    """)
    md("""
    개인 공개 연습 저장소의 루트에 `exercise/checkout.py`, `exercise/test_checkout.py`를 두고, 위 파일을 `.github/workflows/day4-checkout-tests.yml`로 직접 복사하면 해당 저장소의 Actions에서 테스트를 실행할 수 있습니다. 수업 저장소의 자동화 설정은 바꾸지 않습니다.

    선택 읽기 실습은 공개 PR 정보만 받습니다. 비공개 저장소의 학생·고객 자료를 입력하지 않습니다.
    """)
    code("""
    if ALLOW_PUBLIC_GITHUB_READ:
        PUBLIC_REPO = 'YOUR_NAME/YOUR_REPO'  # 본인 공개 연습 저장소로 변경
        PUBLIC_PR = 1
        public_snapshot = fetch_public_pr(PUBLIC_REPO, PUBLIC_PR, allow_network=True)
        print('읽기 완료:', public_snapshot['repo'], public_snapshot['number'])
    else:
        print('GitHub 네트워크 읽기 미실행. 필수 실습은 fixture로 완료됩니다.')
    """)
    code("""
    checks = subprocess.run([sys.executable, '-m', 'pytest', '-q', 'tests/test_day4_pr_review_lab.py'],
                            cwd=ROOT, capture_output=True, text=True, timeout=60)
    print(checks.stdout[-1500:])
    assert checks.returncode == 0, 'PR 계약 테스트를 확인하세요.'
    """)
    md("""
    ### Codex 작업

    > 내 CI YAML에서 비밀키가 필요한 단계가 있는지, PR 제목을 shell 코드로 삽입하는 부분이 있는지, 테스트와 게시 권한이 섞여 있는지 리뷰해 줘. 로컬 테스트와 같은 작업 폴더에서 실행되는지도 확인해 줘. GitHub에 파일을 올리기 전 diff만 먼저 보여 줘.
    """)

    md("""
    ## 4차시 WBS와 Gantt Chart

    WBS는 작업·담당 Role·기간·진행률·선행 작업을 연결한 계획표입니다. AI는 작업을 나누는 초안을 도울 수 있지만 날짜 계산과 진행 상태는 명시적인 규칙과 함수로 검증합니다.
    """)
    code("""
    from labs.day4.office_lab.sheets.student_excel import (
        load_sample, calculate_wbs, validate_wbs, grade_exam,
        write_wbs_copy, write_exam_copy, write_grade_csv,
    )
    tasks = load_sample('wbs_tasks.json')
    AS_OF = '2026-09-25'  # 수업용 프로젝트의 고정 기준일. 오늘 날짜가 아닙니다.
    show_rows(tasks[:5], ['id', 'task', 'role', 'start', 'end', 'progress', 'predecessor'])
    """)
    code("""
    def my_working_days(start, end):
        # 직접 수정 구간: 월~금만 계산. 공휴일 데이터는 아직 반영하지 않습니다.
        first, last = date.fromisoformat(start), date.fromisoformat(end)
        if last < first:
            raise ValueError('END_BEFORE_START')
        return sum((first + timedelta(days=i)).weekday() < 5 for i in range((last-first).days + 1))

    assert my_working_days('2026-09-11', '2026-09-14') == 2  # 금요일·월요일
    assert my_working_days('2026-09-12', '2026-09-13') == 0  # 주말
    expected_failure('종료일 역전', lambda: my_working_days('2026-09-14', '2026-09-11'), 'END_BEFORE_START')
    """)
    md("""
    Excel의 `NETWORKDAYS`와 Python 계산 결과를 비교합니다. 기준일을 고정하면 같은 입력으로 같은 수업 결과를 얻습니다. 공휴일을 제외하려면 공휴일 목록과 적용 기준을 함께 추가해야 합니다.
    """)
    code("""
    wbs_report = calculate_wbs(tasks, AS_OF)
    show_rows(wbs_report, ['id', 'task', 'working_days', 'progress', 'status'])
    for row in wbs_report:
        assert row['working_days'] == my_working_days(row['start'], row['end'])
    """)
    code("""
    broken_tasks = deepcopy(tasks)
    broken_tasks[0]['end'] = '2026-09-01'
    show_rows(validate_wbs(broken_tasks))
    expected_failure('WBS 저장 전 입력 검사', lambda: calculate_wbs(broken_tasks, AS_OF), 'WBS_INPUT_INVALID')
    """)
    code("""
    # 직접 수정 구간: 담당 Role과 완료율을 바꾸고 결과를 확인합니다.
    tasks_my = deepcopy(tasks)
    tasks_my[3]['progress'] = 1
    tasks_my[3]['role'] = '개발 담당'
    show_rows([calculate_wbs(tasks_my, AS_OF)[3]], ['id', 'task', 'role', 'progress', 'status'])
    """)
    code("""
    wbs_path = write_wbs_copy(ROOT, tasks_my, AS_OF, str((RUN/'excel/WBS_my_project.xlsx').relative_to(ROOT)))
    artifacts['내 WBS Excel'] = wbs_path
    file_link(wbs_path)
    from openpyxl import load_workbook
    workbook = load_workbook(wbs_path, data_only=False)
    sheet = workbook['WBS']
    formulas = [{'셀': cell.coordinate, '수식': cell.value} for row in sheet for cell in row if cell.data_type == 'f']
    show_rows(formulas[:6])
    print('조건부서식 범위 수:', len(sheet.conditional_formatting))
    workbook.close()
    """)
    md("""
    ### Codex 작업

    > 내 WBS의 작업 목록과 선행 관계를 읽고 누락된 검증·문서화 작업을 제안해 줘. 날짜는 임의로 바꾸지 마. 내가 선택한 작업의 진행률을 수정하는 코드와 테스트부터 제시해 줘. Excel 수식·조건부서식·입력 셀을 보존하고 새 파일로 저장해 줘.

    저장한 Excel은 Excel 또는 LibreOffice에서 열어 재계산합니다. `openpyxl`은 수식 계산 엔진이 아니며, Python 표는 독립 계산 결과입니다.
    """)

    md("""
    ## 5차시 객관식 채점과 조건부서식

    40문항·28명의 합성 답안을 사용합니다. 정답·오답뿐 아니라 미응답, 결시, 잘못된 보기 번호를 구분합니다. 정답이 누락된 문항을 조용히 0점 처리하면 잘못된 성적을 만들 수 있습니다.
    """)
    code("""
    exam = load_sample('exam_sample.json')
    grading = grade_exam(exam)
    show_rows(grading['students'][-6:])
    print('통계에 포함한 응시자:', grading['eligible_count'])
    """)
    code("""
    def my_answer_status(answer, correct):
        # 직접 수정 구간: 미응답과 입력 오류의 의미를 분리합니다.
        valid = lambda value: not isinstance(value, bool) and isinstance(value, (int, float)) and value in (1,2,3,4)
        if not valid(correct):
            return '정답 확인'
        if answer is None or answer == '':
            return '미응답'
        if not valid(answer):
            return '입력 오류'
        return '정답' if answer == correct else '오답'

    cases = [(1,1,'정답'), (2,1,'오답'), (None,1,'미응답'), (9,1,'입력 오류'), (1,None,'정답 확인')]
    for answer, correct, expected in cases:
        assert my_answer_status(answer, correct) == expected
    show_rows([{'입력': a, '정답': k, '판정': my_answer_status(a,k)} for a,k,_ in cases])
    """)
    md("""
    ### 정답표와 통계의 기준

    보기 9는 오답이 아니라 입력 오류입니다. 원본 답안 확인 후 수정합니다. 실습에서는 의도적으로 넣은 보기 9를 1로 고칩니다. 실제 시험에서는 근거 없이 보기를 추정하지 않습니다.
    """)
    code("""
    exam_my = deepcopy(exam)
    exam_my['students'][25]['answers'][8] = 1
    grading_my = grade_exam(exam_my)
    show_rows([grading['students'][25], grading_my['students'][25]])
    assert grading['students'][25]['score'] is None
    assert grading_my['students'][25]['score'] is not None
    """)
    code("""
    key_missing = deepcopy(exam_my)
    key_missing['key'][0] = None
    missing_result = grade_exam(key_missing)
    assert missing_result['students'][0]['score'] is None
    print('정답 누락 상태:', missing_result['students'][0]['status'])
    key_changed = deepcopy(exam_my)
    key_changed['key'][0] = 2
    changed_result = grade_exam(key_changed)
    show_rows([{'버전': '기존 정답표', '평균': grading_my['mean']},
               {'버전': '1번 정답 변경', '평균': changed_result['mean']}])
    """)
    code("""
    exam_path = write_exam_copy(ROOT, exam_my, str((RUN/'excel/Exam_my_class.xlsx').relative_to(ROOT)))
    changed_key_path = write_exam_copy(ROOT, key_changed, str((RUN/'excel/Exam_changed_key.xlsx').relative_to(ROOT)))
    csv_path = write_grade_csv(ROOT, exam_my, str((RUN/'excel/grading_report.csv').relative_to(ROOT)))
    artifacts['내 채점 Excel'], artifacts['채점 CSV'] = exam_path, csv_path
    artifacts['정답 변경 비교 Excel'] = changed_key_path
    file_link(exam_path)
    file_link(changed_key_path)
    file_link(csv_path)
    print(csv_path.read_text(encoding='utf-8-sig').splitlines()[0])
    """)
    code("""
    workbook = load_workbook(exam_path, data_only=False)
    sheet = workbook['채점']
    rules = [{'적용 범위': str(rule), '규칙 수': len(sheet.conditional_formatting[rule])} for rule in sheet.conditional_formatting]
    show_rows(rules[:8])
    print('정답 셀 D8:', sheet['D8'].value, '/ 첫 답안 셀 D10:', sheet['D10'].value)
    workbook.close()
    changed_workbook = load_workbook(changed_key_path, data_only=False)
    assert changed_workbook['채점']['D8'].value == 2
    print('정답 변경 비교 파일 D8:', changed_workbook['채점']['D8'].value)
    changed_workbook.close()
    show_rows(grading_my['questions'][:5])
    """)
    md("""
    ### Codex 작업

    > 합성 답안으로 만든 채점 Excel을 검토해 줘. 정답 행과 학생 답안의 비교 참조가 고정되어 있는지, 결시·미응답·보기 오류를 구분하는지 확인해 줘. 정답 1개를 바꾼 뒤 점수·등수·문항별 정답률이 함께 바뀌는 테스트를 추가해 줘. 실제 학생 점수나 원본 Google Sheets는 수정하지 마.
    """)

    md("""
    ## 6차시 Word 이력서 개선

    AI가 경력을 더 읽기 쉽게 정리하도록 하되 없던 성과·직함·수치를 만들지 않도록 합니다. 원본 사실에 ID를 붙여 변경 문장과 연결하고, 사실 검증이 필요한 부분은 사람의 확인 항목으로 남깁니다.
    """)
    code("""
    from labs.day4.office_lab.documents.resume_lab import read_fact_map, audit_rewrite, build_resume, DocumentLabError
    FACT_SOURCE = 'labs/day4/office_lab/documents/career_facts.json'
    facts = read_fact_map(ROOT, FACT_SOURCE)
    fact_rows = [{'id': fact['id'], '산업': role['domain'], '원문': fact['text']} for role in facts['roles'] for fact in role['facts']]
    show_rows(fact_rows)
    """)
    code("""
    # 직접 수정 구간: 의미와 수치를 유지한 문장 개선만 입력합니다.
    my_rewrites = {'medical-polyp': '대장내시경 영상 기반 용종 탐지 및 진단 서비스 개발과 운영'}
    show_rows([{'사실 ID': key, '개선 문장': value} for key, value in my_rewrites.items()])
    """)
    code("""
    rewrite_check = audit_rewrite(facts, my_rewrites)
    show_rows([{'수치 검사': rewrite_check['numeric_guard'], '의미 사실 검증': rewrite_check['semantic_claims_verified'], '상태': rewrite_check['status']}])
    assert rewrite_check['status'] == 'REQUIRES_HUMAN_REVIEW'
    """)
    code("""
    expected_failure('근거 없는 수치', lambda: audit_rewrite(facts,
        {'education-latency': '평가 시간을 1초로 단축'}), 'UNSUPPORTED_NUMERIC_CLAIM')
    expected_failure('없는 경력 ID', lambda: audit_rewrite(facts,
        {'unknown-career': '새로운 경력'}), 'UNKNOWN_FACT_REFERENCE')
    """)
    code("""
    before = build_resume(ROOT, FACT_SOURCE, RUN/'word/Resume_Before.docx', variant='before')
    after = build_resume(ROOT, FACT_SOURCE, RUN/'word/Resume_After.docx', proposals=my_rewrites, variant='after')
    artifacts['이력서 수정 전'], artifacts['이력서 수정 후'] = Path(before['path']), Path(after['path'])
    file_link(before['path'])
    file_link(after['path'])
    """)
    code("""
    from docx import Document
    document = Document(after['path'])
    paragraphs = [{'스타일': p.style.name, '내용': p.text} for p in document.paragraphs if p.text]
    show_rows(paragraphs[:9])
    assert document.paragraphs[0].style.name == 'Title'
    assert any(p.style.name == 'Heading 1' for p in document.paragraphs)
    """)
    md("""
    제목·소제목·본문은 Word 스타일로 구분합니다. 문장을 짧게 바꾸는 것과 글자 크기를 일괄 변경하는 것은 다른 작업입니다. 생성 후 Word에서 모든 페이지를 열어 줄바꿈·글꼴·페이지 나눔을 확인합니다. 숫자 검사는 새 수치를 찾아낼 뿐, 기존 내용의 진실성까지 검증하지 않습니다.
    """)
    md("""
    ### Codex 작업

    > 사실 목록과 수정 전 이력서를 기준으로, 현재 Role과 산업별 경험이 먼저 보이도록 문장과 순서를 개선해 줘. 직함·수치·성과를 새로 만들지 마. 근거가 없는 내용은 확인 질문으로 남겨 줘. 원본은 보존하고 새 Word 파일로 만들고, 수정한 문장과 사실 ID의 대응표를 함께 보여 줘.
    """)

    md("""
    ## 7차시 데이터 기반 PPT 생성

    같은 WBS 입력으로 일정 현황과 보고서 표를 만듭니다. 숫자는 코드가 계산하고, AI에는 독자에게 필요한 설명·순서·표현을 요청합니다. PPT를 만든 뒤에도 수치·페이지·가독성 검사가 필요합니다.
    """)
    code("""
    from labs.day4.office_lab.presentations.brief_data import prepare_brief
    brief = prepare_brief(ROOT, as_of=AS_OF, tasks=tasks_my)
    show_rows([{'페이지': i+1, '제목': slide['title'].replace('\\n', ' ')} for i, slide in enumerate(brief['slides'])])
    """)
    md("""
    ### 직접 수정 구간

    표의 숫자는 계산 결과를 유지합니다. 제목·부제·본문은 수신자의 업무에 맞게 바꿀 수 있습니다. 4차시에서 바꾼 작업 완료율이 이번 보고서의 완료 업무 수와 같은지 확인합니다.
    """)
    code("""
    brief['slides'][0]['subtitle'] = '수업용 프로젝트 진행 현황  ' + AS_OF
    def my_brief_validation(data):
        if data.get('schema_version') != 1 or len(data.get('slides', [])) != 7:
            raise ValueError('BRIEF_SCHEMA_INVALID')
        for slide in data['slides']:
            if not isinstance(slide.get('title'), str) or not slide['title'].strip():
                raise ValueError('BRIEF_TITLE_MISSING')
        return 'PASS'
    assert my_brief_validation(brief) == 'PASS'
    broken_brief = deepcopy(brief)
    broken_brief['slides'] = broken_brief['slides'][:6]
    expected_failure('보고서 페이지 누락', lambda: my_brief_validation(broken_brief), 'BRIEF_SCHEMA_INVALID')
    complete = sum(row['status'] == '완료' for row in calculate_wbs(tasks_my, AS_OF))
    assert int(brief['slides'][1]['table'][1][1]) == complete
    print('WBS와 PPT의 완료 업무 수:', complete)
    """)
    code("""
    brief_path = RUN/'ppt/project_brief.json'
    brief_path.parent.mkdir(parents=True, exist_ok=True)
    brief_path.write_text(json.dumps(brief, ensure_ascii=False, indent=2), encoding='utf-8')
    file_link(brief_path)
    """)
    md("""
    ### Node.js 준비

    실제 PPTX 생성에는 Node.js와 공개 라이브러리 PptxGenJS가 필요합니다. Node.js 설치 후 저장소 루트 터미널에서 아래 명령을 한 번 실행합니다.

    ```bash
    npm install --prefix labs/day4/office_lab/presentations
    ```

    설치를 자동으로 몰래 실행하지 않습니다. 준비되지 않았으면 아래 셀은 필요한 명령을 보여주며, 입력·검증 단계와 제공한 PPT 예시는 계속 확인할 수 있습니다. 실행본에는 실제 PPTX 생성 여부를 따로 표시합니다.
    """)
    code("""
    NODE = shutil.which('node')
    PPT_PACKAGE = ROOT/'labs/day4/office_lab/presentations'
    PPT_NODE_READY = bool(NODE and (PPT_PACKAGE/'node_modules/pptxgenjs/package.json').is_file())
    print('Node.js:', '준비됨' if NODE else '설치 필요')
    print('PptxGenJS:', '준비됨' if PPT_NODE_READY else '위 npm install 명령 실행 필요')
    """)
    code("""
    ppt_path = RUN/'ppt/Project_Brief_my_project.pptx'
    if PPT_NODE_READY:
        ppt_run = subprocess.run([NODE, str(PPT_PACKAGE/'student_ppt.mjs'), str(brief_path), str(ppt_path)],
                                 cwd=ROOT, capture_output=True, text=True, timeout=60)
        if ppt_run.returncode:
            raise RuntimeError('PPT 생성 실패. 선택한 JSON·Node.js·PptxGenJS 설치를 확인하세요.')
        artifacts['내 보고서 PPT'] = ppt_path
        print('PPT 생성 완료:', ppt_path.name)
        file_link(ppt_path)
    else:
        print('PPT 파일 생성 미실행: 설치 후 이 셀부터 다시 실행하세요.')
    """)
    code("""
    from zipfile import ZipFile
    from xml.etree import ElementTree as ET
    if ppt_path.is_file():
        with ZipFile(ppt_path) as package:
            pages = [name for name in package.namelist() if name.startswith('ppt/slides/slide') and name.endswith('.xml')]
            assert len(pages) == 7
            texts = [text for name in pages for node in ET.fromstring(package.read(name)).iter() if node.tag.endswith('}t') for text in [node.text or '']]
            assert '일정 현황' in texts
        print('PPT 구조 검사: 7페이지, 일정 현황 제목 확인')
        print('다음 확인: PowerPoint에서 모든 페이지의 줄바꿈·표·글꼴을 직접 확인')
    """)
    md("""
    ### Codex 작업

    > 내 WBS 입력과 project_brief.json을 기준으로 보고서 PPT를 개선해 줘. 표 수치와 상태는 계산 결과를 유지하고, 제목은 짧은 명사형으로 정리해 줘. 반복되는 설명을 줄이고 표와 비교 중심으로 구성해 줘. 생성 후 PPT를 열어 작은 글씨·잘린 표·페이지 밖 요소를 확인할 체크리스트도 작성해 줘.
    """)

    md("""
    ## 8차시 결과 연결과 미니 프로젝트 준비

    오늘 만든 결과를 한 곳에서 열고, 다음 주에 확장할 서비스를 고릅니다. 여러 파일을 만든 것에서 끝내지 않고 입력 변경·재실행·오류 처리·사람 확인까지 이어지는 작은 도구를 목표로 합니다.
    """)
    code("""
    def my_output_manifest(named_paths):
        rows = []
        for label, path in named_paths.items():
            path = Path(path).resolve()
            if not path.is_relative_to(RUN):
                raise ValueError('PATH_OUTSIDE_RUN')
            rows.append({'자료': label, '파일': path.name, '존재': path.is_file(),
                         '크기KB': round(path.stat().st_size/1024, 1) if path.is_file() else 0})
        return rows
    manifest = my_output_manifest(artifacts)
    show_rows(manifest)
    assert all(item['존재'] for item in manifest)
    expected_failure('다른 폴더의 파일 혼입', lambda: my_output_manifest({'범위 밖': ROOT/'README.md'}), 'PATH_OUTSIDE_RUN')
    """)
    code("""
    expected_outputs = ['PR 리뷰 화면', '리뷰 Markdown', '내 WBS Excel', '내 채점 Excel', '채점 CSV', '이력서 수정 후']
    assert all(label in artifacts for label in expected_outputs)
    print('Python 기반 필수 결과:', len(expected_outputs), '개 확인')
    print('PPT 실제 생성:', '완료' if '내 보고서 PPT' in artifacts else 'Node.js 설치 후 실행 필요')
    print('내 코드 테스트:', current_test['status'], '/ 참고 구현 테스트:', reference_test['status'])
    """)
    code("""
    from urllib.parse import quote
    links = ''.join('<li><a href="' + quote(str(Path(path).relative_to(RUN)).replace(os.sep, '/')) + '">' + html.escape(label) + '</a></li>' for label,path in artifacts.items())
    index_path = RUN/'index.html'
    index_path.write_text('<!doctype html><html lang="ko"><meta charset="utf-8"><title>4주차 내 실습 결과</title><style>body{font:20px/1.7 sans-serif;max-width:900px;margin:60px auto;padding:20px;color:#111}a{color:#111}li{margin:14px 0}</style><h1>4주차 내 실습 결과</h1><p>이번 실행에서 생성한 로컬 파일입니다. 원격 게시·메일 발송 없음.</p><ul>' + links + '</ul></html>', encoding='utf-8')
    file_link(index_path)
    print('탐색기 또는 Finder에서 index.html을 열어 결과를 확인하세요.')
    """)
    md("""
    ### 다음 주 마지막 3시간

    | 선택 과제 | 입력 | 직접 만들 기능 | 개선 확인 |
    |---|---|---|---|
    | 프로젝트 일정 도우미 | 작업 목록 | WBS·날짜 검사·Gantt Excel | 선행 작업 충돌·진행률 변경 |
    | 채점 도우미 | 합성 정답·답안 | Excel·CSV·입력 오류 표시 | 결시·정답 변경·등수 재계산 |
    | 경력 문서 도우미 | 본인이 공유 가능한 경력 | Word 개선·사실 대응표 | 근거 없는 수치·원본 보존 |
    | PR 검토 도우미 | 본인 연습 PR | 리뷰·테스트·게시 전 확인 | 새 commit·중복·권한 |

    3시간 중 150분은 제작, 마지막 30분은 쉬는 시간·Q&A입니다. 강제 발표는 없으며 수정 전·후 화면과 테스트 결과를 개인 폴더에 남깁니다.

    | 시간 | 배분 | 진행 내용 | 확인할 결과 |
    |---|---:|---|---|
    | 15:00–15:20 | 20분 | 사용자·입력·완료 기준 설계 | 기능 하나, 정상 입력 1건, 실패 입력 2건 |
    | 15:20–16:00 | 40분 | 기본 코드 실행과 함수·입력 변경 | 자신의 결과 파일 첫 생성 |
    | 16:00–16:40 | 40분 | 실패 재현과 코드·테스트 개선 | 예상 오류와 복구 결과 |
    | 16:40–17:10 | 30분 | 결과·서식·화면·실행 안내 개선 | 결과 파일과 README |
    | 17:10–17:30 | 20분 | 새 폴더 재실행과 개선 기록 | 재실행 결과와 개선 전후 비교 |
    | 17:30–18:00 | 30분 | 쉬는 시간·Q&A | 선택 질문과 실습 복구 |
    """)
    md("""
    ### ChatGPT Desktop·Claude Desktop 대화 예시

    **첫 요청**

    > 나는 [누가 쓰는지]를 위한 [작은 도구]를 만들고 싶어. 입력은 [공개 합성 파일], 출력은 [xlsx/docx/pptx/웹 화면]이야. 먼저 파일 구조와 실패 조건을 정리해 줘. 필요한 계정·설치·외부 전송은 실행 전에 알려 줘. 원본은 수정하지 말고 결과 폴더에 새 파일로 만들어 줘.

    **구현 요청**

    > 먼저 정상 입력 한 개가 끝까지 작동하게 만들어 줘. 그런 다음 빈 입력·잘못된 값·중복 실행 테스트를 추가해 줘. 내가 수정할 함수와 실행 명령을 표시하고 실제 테스트 결과를 보여 줘.

    **개선 요청**

    > 결과 파일을 열어서 확인했어. [문제]를 고쳐 줘. 기존 수치와 파일 경로는 유지하고, 바뀐 코드와 테스트를 먼저 보여 줘. 외부 게시나 이메일 발송은 하지 마.

    앱·계정마다 사용할 수 있는 파일 도구와 연결 기능이 다릅니다. 기능이 없으면 오늘 제공한 로컬 코드와 양식을 실행해 같은 결과 파일을 만들 수 있습니다.
    """)
    code("""
    next_task = RUN/'next_project_task.md'
    next_task.write_text('# 다음 주 프로젝트 입력\\n\\n- 사용자:\\n- 반복 업무:\\n- 입력 파일:\\n- 결과 파일 또는 화면:\\n- 정상 입력 예시:\\n- 실패 입력 예시:\\n- 사람 확인 단계:\\n- 외부 연결 필요 여부:\\n- 개선 전후 비교 방법:\\n', encoding='utf-8')
    file_link(next_task)
    """)
    md("""
    ### 개인 점검

    내 코드의 실패 원인을 설명할 수 있는지, 입력을 바꿨을 때 결과가 다시 계산되는지, 문서에 근거 없는 사실이 추가되지 않았는지 확인합니다. 실행이 안 되는 부분은 오류 코드·사용한 명령·파일 위치를 함께 기록하면 강사가 빠르게 도울 수 있습니다.

    전체 결과는 위에 출력된 `output/day4-notebook-runs/...`에 있습니다. 다시 Run All을 하면 새 폴더를 만들고 이전 결과를 보존합니다.
    """)
    book = nbf.v4.new_notebook(cells=cells, metadata={
        'kernelspec': {'display_name': 'Python 3', 'language': 'python', 'name': 'python3'},
        'language_info': {'name': 'python', 'version': '3.12'},
        'day4': {'network_default': False, 'automatic_publish': False, 'automatic_student_fix': False},
    })
    for index, cell in enumerate(book.cells):
        cell.id = f'day4-{index+1:03d}'
        if cell.cell_type == 'code':
            compile(cell.source, f'<day4-cell-{index+1}>', 'exec')
    nbf.validate(book)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    nbf.write(book, OUTPUT)
    print(f'{OUTPUT}: {len(cells)} cells, {sum(c.cell_type == "code" for c in cells)} code cells')
    return book


def execute(book):
    """Use this exact Python in a temporary kernel without changing user settings."""
    from nbclient import NotebookClient
    from jupyter_client import KernelManager
    from jupyter_client.kernelspec import KernelSpecManager
    with tempfile.TemporaryDirectory(prefix='day4-notebook-kernel-') as temp:
        kernel_dir = Path(temp) / 'day4-local'
        kernel_dir.mkdir()
        kernel = {'argv': [sys.executable, '-m', 'ipykernel_launcher', '-f', '{connection_file}'],
                  'display_name': 'Day4 local validation', 'language': 'python'}
        (kernel_dir/'kernel.json').write_text(json.dumps(kernel), encoding='utf-8')
        manager = KernelManager(kernel_name='day4-local', kernel_spec_manager=KernelSpecManager(kernel_dirs=[temp]))
        client = NotebookClient(book, km=manager, timeout=180, resources={'metadata': {'path': str(ROOT)}})
        try:
            client.execute()
        finally:
            if manager.has_kernel:
                manager.shutdown_kernel(now=True)
        destination = OUTPUT.with_name(OUTPUT.stem + '.executed.ipynb')
        nbf.write(book, destination)
        print(f'Executed: {destination}')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--execute', action='store_true')
    args = parser.parse_args()
    book = build()
    if args.execute:
        execute(book)
